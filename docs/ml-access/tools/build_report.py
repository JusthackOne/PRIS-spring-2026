"""Create REPORT_4.docx from the single Markdown master document."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from PIL import Image

REPO=Path(__file__).resolve().parents[3]
BASE=REPO/'docs/ml-access'
MASTER=BASE/'ML_SYSTEM_DESIGN.md'
DOC=Document()
FIGURES=[]
NORMAL=DOC.styles['Normal']
NORMAL.font.name='Arial'; NORMAL.font.size=Pt(11)
NORMAL.paragraph_format.space_after=Pt(7)
NORMAL.paragraph_format.line_spacing=1.15
for s,size in [('Title',24),('Heading 1',17),('Heading 2',14),('Heading 3',12)]:
    st=DOC.styles[s];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.space_before=Pt(14);st.paragraph_format.space_after=Pt(8)
    st.paragraph_format.keep_with_next=True
DOC.styles['Title'].font.bold=False
DOC.styles['Caption'].font.name='Arial';DOC.styles['Caption'].font.size=Pt(10)
DOC.styles['Caption'].font.color.rgb=RGBColor(60,60,60)
DOC.core_properties.author='Недиков Михаил Олегович'
DOC.core_properties.title='Кейс №2 Идентификация сотрудников по лицу на проходных предприятия'
DOC.core_properties.subject='Проектирование информационной системы с ML'
DOC.core_properties.keywords='MLSD, IDEF0, UML, идентификация, СКУД'

def configure(sec, landscape=False, a3=False):
    sec.orientation=WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    short,long=(11.69,16.54) if a3 else (8.27,11.69)
    sec.page_width=Inches(long if landscape else short)
    sec.page_height=Inches(short if landscape else long)
    sec.left_margin=sec.right_margin=Inches(.63)
    sec.top_margin=sec.bottom_margin=Inches(.62)
    sec.header_distance=sec.footer_distance=Inches(.25)
configure(DOC.sections[0])

footer=DOC.sections[0].footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('Контур доступа  |  Недиков М. О.  |  ');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
header=DOC.sections[0].header.paragraphs[0]
r=header.add_run('Проектное задание  •  Кейс №2  •  ML System Design');r.font.name='Arial';r.font.size=Pt(8);r.font.color.rgb=RGBColor(100,100,100)

def url_target(target):
    if target.startswith('http'):return target
    if target.startswith('#'):return 'https://github.com/JusthackOne/PRIS-spring-2026/blob/project/ml-access-case2/docs/ml-access/ML_SYSTEM_DESIGN.md'+target
    return 'https://github.com/JusthackOne/PRIS-spring-2026/blob/project/ml-access-case2/docs/ml-access/'+target

def inline(p,text):
    pattern=r'(\[[^\]]+\]\([^)]+\)|\*\*.*?\*\*|`[^`]+`)'
    for token in re.split(pattern,text):
        if not token:continue
        m=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',token)
        if m:
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(url_target(m[2]),RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'23567C');pr.append(col);r.append(pr)
            t=OxmlElement('w:t');t.text=m[1];r.append(t);h.append(r);p._p.append(h)
        elif token.startswith('**'):
            p.add_run(token[2:-2]).bold=True
        elif token.startswith('`'):
            rr=p.add_run(token[1:-1]);rr.font.name='Consolas';rr.font.size=Pt(9)
        else:p.add_run(token)

def table(rows):
    t=DOC.add_table(rows=1,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    n=len(rows[0]); ratios={2:[.30,.70],3:[.25,.45,.30],4:[.25,.33,.25,.17]}[n]
    avail=7.01
    for c,ratio in zip(t.columns,ratios):c.width=Inches(avail*ratio)
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,cell in enumerate(cells):
            cell.width=Inches(avail*ratios[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1.05
            inline(p,row[j])
            for rr in p.runs:rr.font.size=Pt(9.5);rr.bold=i==0
            tcpr=cell._tc.get_or_add_tcPr()
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EEF4' if i==0 else ('F7F9FB' if i%2==0 else 'FFFFFF'));tcpr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'80');el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcpr.append(borders)
        trpr=t.rows[i]._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trpr.append(cant)
        if i==0:trpr.append(OxmlElement('w:tblHeader'))
    DOC.add_paragraph().paragraph_format.space_after=Pt(0)

def figure(alt,path):
    pth=BASE/path;w,h=Image.open(pth).size
    a3=pth.stem in ['er','mvp','sequence','architecture','components']
    wide=w/h>1.1
    sec=DOC.add_section(WD_SECTION_START.NEW_PAGE);configure(sec,wide,a3)
    p=DOC.add_paragraph('Приложение А  |  '+alt,style='Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    # Reserve caption and section-break paragraph; do not force a blank trailing page.
    aw=(15.28 if wide else 10.43) if a3 else (10.43 if wide else 7.01)
    ah=(9.50 if wide else 14.35) if a3 else (6.08 if wide else 9.30)
    scale=min(aw/w,ah/h)
    p=DOC.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(0)
    shape=p.add_run().add_picture(str(pth),width=Inches(w*scale),height=Inches(h*scale))
    shape._inline.docPr.set('descr',alt)

lines=MASTER.read_text(encoding='utf-8').splitlines();i=0
while i<len(lines):
    line=lines[i].strip()
    if not line:i+=1;continue
    if line.startswith('**Навигация:'):i+=1;continue
    im=re.match(r'!\[([^]]+)\]\(([^)]+)\)',line)
    if im:
        FIGURES.append((im[1],im[2]))
        p=DOC.add_paragraph(im[1]+' (приложение А).',style='Caption')
        i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            parts=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[-: ]+',x) for x in parts):rows.append(parts)
            i+=1
        table(rows);continue
    if line.startswith('#'):
        depth=len(line)-len(line.lstrip('#'));title=line.lstrip('#').strip()
        p=DOC.add_paragraph(style='Title' if depth==1 else 'Heading '+str(min(depth-1,3)))
        inline(p,title);i+=1;continue
    if line.startswith('- '):
        p=DOC.add_paragraph(style='List Bullet');inline(p,line[2:]);i+=1;continue
    p=DOC.add_paragraph();inline(p,line)
    i+=1

for alt,path in FIGURES:figure(alt,path)
# Remove inherited template title rules, theme borders and first-page switches.
for tree in [DOC._element,DOC.styles.element]:
    for el in list(tree.iter(qn('w:pBdr'))):el.getparent().remove(el)
for sec in DOC.sections:sec.different_first_page_header_footer=False
DOC.save(REPO/'REPORT_4.docx')
(BASE/'MASTER.me').write_bytes(MASTER.read_bytes())
print(REPO/'REPORT_4.docx')
